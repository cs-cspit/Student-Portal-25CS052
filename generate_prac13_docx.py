import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Page Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Header / College Info
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("CHANDUBHAI S. PATEL INSTITUTE OF TECHNOLOGY (CSPIT)\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    run_inst.bold = True
    run_inst.font.size = Pt(13)
    run_inst.font.color.rgb = RGBColor(30, 41, 59)

    run_sub = title_p.add_run("ITUE203 - Web Development Frameworks / Database Systems\n")
    run_sub.bold = True
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(37, 99, 235)

    run_prac = title_p.add_run("PRACTICAL - 13 LABORATORY REPORT\n")
    run_prac.bold = True
    run_prac.font.size = Pt(15)
    run_prac.font.color.rgb = RGBColor(15, 23, 42)

    run_theme = title_p.add_run("Practical 13: PHP JSON API and AJAX/Fetch-Based CRUD without Page Reload\n")
    run_theme.italic = True
    run_theme.font.size = Pt(11)
    run_theme.bold = True

    # Student metadata table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_table.rows[0].cells[0].text = "Student Name: Jay Navadiya"
    meta_table.rows[0].cells[1].text = "Student ID / Roll No: 25CS052"
    meta_table.rows[1].cells[0].text = "Degree / Semester: B.Tech CSE (Div-I)"
    meta_table.rows[1].cells[1].text = "Date of Submission: October 21, 2026"

    for row in meta_table.rows:
        for cell in row.cells:
            set_cell_background(cell, "F1F5F9")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.bold = True
                    run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Aim & Objectives
    h1 = doc.add_heading("1. Aim, Objectives & Problem Definition", level=1)
    h1.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    p_aim = doc.add_paragraph()
    p_aim.add_run("AIM: ").bold = True
    p_aim.add_run(
        "To build robust, RESTful PHP JSON API endpoints for student management, consume them asynchronously using the modern JavaScript "
        "Fetch API (AJAX), and perform full Create, Read, Update, and Delete (CRUD) operations dynamically without full page reloads. "
        "The system incorporates real-time debounced AJAX search, multi-criteria filtering, multi-layer input validation, centralized error handling, "
        "HTTP status codes (200, 201, 400, 404, 409, 500), and DevTools Network inspection."
    )

    doc.add_heading("1.1 Problem Definition", level=2)
    p_def = doc.add_paragraph()
    p_def.add_run(
        "Create simple PHP JSON API endpoints for events or students and consume them using Fetch API. "
        "Perform list, add, update, and delete operations asynchronously without full page reload."
    )

    doc.add_heading("1.2 Course Outcome (CO) Mapping", level=2)
    p_co = doc.add_paragraph()
    p_co.add_run("• CO3: ").bold = True
    p_co.add_run("Design responsive, asynchronous, and accessible user interfaces using HTML5, modern CSS, and JavaScript DOM manipulation.\n")
    p_co.add_run("• CO4: ").bold = True
    p_co.add_run("Implement server-side RESTful API architectures using PHP, JSON data interchange, and secure MySQLi prepared statements.\n")
    p_co.add_run("• CO5: ").bold = True
    p_co.add_run("Integrate frontend client-side Fetch API with backend database endpoints, handling HTTP status codes, loading indicators, and centralized error handling.")

    doc.add_heading("1.3 Learning Outcomes", level=2)
    lo_items = [
        ("Asynchronous Frontend-Backend Integration", "Students will master connecting modern browser Fetch API requests with server-side PHP endpoints without full page reloads."),
        ("RESTful JSON API Design", "Students will understand standard REST conventions, JSON content serialization, and HTTP status codes (200 OK, 201 Created, 400 Bad Request, 404 Not Found, 409 Conflict, 500 Internal Error)."),
        ("Single-Page Application (SPA) UX Patterns", "Students will implement interactive UI states including asynchronous spinners, skeleton loaders, success toast notifications, and client/server validation alerts."),
        ("Real-time AJAX Debouncing", "Students will implement intermediate debounce algorithms (300ms) to throttle rapid search keystrokes and prevent excessive backend server strain."),
        ("Browser Network Diagnostics", "Students will inspect and interpret HTTP requests, response headers, payload streams, and timing metrics within the browser Developer Tools Network tab.")
    ]
    for title, desc in lo_items:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.add_run(f"{title}: ").bold = True
        p_item.add_run(desc)

    # Section 2: Key Questions / Analysis & Interpretation
    h2 = doc.add_heading("2. Key Questions & Technical Analysis", level=1)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    qa_list = [
        ("1. Do PHP endpoints return valid JSON responses?",
         "Yes. In api.php, the server explicitly transmits strict headers: header('Content-Type: application/json; charset=UTF-8') and header('Cache-Control: no-store, no-cache'). All response payloads are assembled into a structured JSON envelope containing 'status', 'code', 'message', 'data', 'timestamp', and optional 'errors', and encoded via json_encode($response, JSON_PRETTY_PRINT | JSON_UNESCAPED_SLASHES)."),

        ("2. Are GET and POST requests handled properly?",
         "Yes. api.php implements a centralized router inspecting $_SERVER['REQUEST_METHOD'] with method override support (via $_POST['_method'], ?action=, or X-HTTP-Method-Override). GET requests retrieve all students or a single record (?id=X), POST requests create new student records (HTTP 201), PUT requests update existing records (HTTP 200), and DELETE requests permanently remove records (HTTP 200)."),

        ("3. Are loading, success, and error states shown on the frontend?",
         "Yes. In index.html and script.js:\n"
         "• Loading State: A pulsing loading spinner overlay displays during asynchronous network round-trips.\n"
         "• Success State: Dismissible animated green toast banners appear immediately when records are created, updated, or deleted.\n"
         "• Error State: Field-level red warning outlines and server-side alert boxes display when duplicate emails or invalid parameters are submitted."),

        ("4. Is input validation performed before database operations?",
         "Yes. Two-stage validation is enforced:\n"
         "1. Client-Side: script.js validates non-empty inputs and regex email syntax before invoking fetch().\n"
         "2. Server-Side: api.php verifies required fields, string lengths, filter_var($email, FILTER_VALIDATE_EMAIL), whitelisted department options, and checks for existing duplicate email addresses using MySQLi prepared statements (returning HTTP 409 Conflict if violated).")
    ]
    for q, a in qa_list:
        doc.add_heading(q, level=2)
        doc.add_paragraph(a)

    # Section 3: Complete Source Code Implementation
    h3 = doc.add_heading("3. Complete Source Code Implementation", level=1)
    h3.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    base_dir = "/Users/jaynavadiya07/Desktop/Project"
    code_files = [
        ("db.php", "Database Connection using MySQLi & PDO", os.path.join(base_dir, "db.php")),
        ("api.php", "RESTful PHP JSON API Endpoints (CRUD)", os.path.join(base_dir, "api.php")),
        ("index.html", "Single-Page Application UI & DevTools Inspector", os.path.join(base_dir, "index.html")),
        ("script.js", "Asynchronous Fetch API Engine & Debounce Search", os.path.join(base_dir, "script.js")),
        ("schema.sql", "MySQL Relational Database Schema & Seed Data", os.path.join(base_dir, "schema.sql"))
    ]

    for fname, fdesc, fpath in code_files:
        doc.add_heading(f"3.{code_files.index((fname, fdesc, fpath))+1} File: {fname} – {fdesc}", level=2)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                content = f.read()
        else:
            content = "// File not found"

        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        set_cell_background(cell, "F8FAFC")
        cell_p = cell.paragraphs[0]
        cell_p.paragraph_format.space_before = Pt(4)
        cell_p.paragraph_format.space_after = Pt(4)
        run = cell_p.add_run(content)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 4: Experimental Screenshots
    h4 = doc.add_heading("4. Experimental Results & Verification Screenshots", level=1)
    h4.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    screenshots_dir = os.path.join(base_dir, "PRAC SS", "Practical 13")
    screenshots = [
        ("01_student_list_interface.png", "UI: Asynchronous Student Catalog Loaded via Fetch GET (api.php)"),
        ("02_add_student_modal.png", "UI: Add New Student Modal Dialog"),
        ("03_add_student_form_filled.png", "UI: Add Student Form Input Data Entry (Pooja Trivedi)"),
        ("04_add_student_success.png", "UI: Successful Asynchronous Record Insertion (HTTP 201 Created) without Page Reload"),
        ("05_edit_student_modal.png", "UI: Edit Student Record Modal Populated with Existing Data"),
        ("06_edit_student_success.png", "UI: Successful Asynchronous Record Update (HTTP 200 OK) without Page Reload"),
        ("07_delete_student_confirmation_modal.png", "UI: Delete Record Confirmation Warning Dialog"),
        ("08_delete_student_success.png", "UI: Successful Asynchronous Deletion (HTTP 200 OK) with Dynamic Row Removal"),
        ("09_loading_state.png", "UI: Asynchronous Fetch Loading Spinner & Progress State"),
        ("10_error_state_client_validation.png", "UI: Client-Side Input Validation Error Indicators"),
        ("11_error_state_duplicate_conflict.png", "UI: Server-Side Duplicate Email Validation Error (HTTP 409 Conflict)"),
        ("12_ajax_search_debounce.png", "Intermediate Extension: Real-Time AJAX Search with 300ms Debounce"),
        ("13_network_tab_overview_all_requests.png", "Browser Network Tab: Overview of All Fetch Requests, HTTP Verbs & Status Codes"),
        ("14_network_tab_fetch_get_response.png", "Browser Network Tab: GET /api.php Response Headers & Structured JSON Response"),
        ("15_network_tab_fetch_post_payload_201.png", "Browser Network Tab: POST /api.php Request Payload & 201 Created Response"),
        ("16_network_tab_fetch_put_payload_200.png", "Browser Network Tab: PUT /api.php Request Payload & 200 OK Response"),
        ("17_network_tab_fetch_delete_200.png", "Browser Network Tab: DELETE /api.php Status 200 OK Response"),
        ("18_network_tab_fetch_error_409_conflict.png", "Browser Network Tab: POST Conflict (409) Error JSON Response Envelope"),
        ("19_raw_api_php_json_endpoint.png", "Direct Browser Endpoint: Standalone JSON Response Output (GET /api.php)")
    ]

    for fname, fcaption in screenshots:
        fpath = os.path.join(screenshots_dir, fname)
        if os.path.exists(fpath):
            doc.add_heading(f"4.{screenshots.index((fname, fcaption))+1} {fcaption}", level=2)
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_after = Pt(4)
            run_img = p_img.add_run()
            run_img.add_picture(fpath, width=Inches(6.5))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            run_cap = p_cap.add_run(f"Figure 4.{screenshots.index((fname, fcaption))+1}: {fcaption} ({fname})")
            run_cap.font.size = Pt(9)
            run_cap.italic = True
            run_cap.font.color.rgb = RGBColor(100, 116, 139)

    # Section 5: Conclusion
    h5 = doc.add_heading("5. Conclusion", level=1)
    h5.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph(
        "In Practical 13, a full-stack asynchronous student management portal was successfully implemented using pure PHP RESTful JSON APIs "
        "and modern JavaScript Fetch API. All CRUD operations (Create, Read, Update, Delete) operate seamlessly in the background without requiring "
        "any browser page reloads. The system achieves robust reliability through dual-stage input validation, standard HTTP status code communication "
        "(200, 201, 400, 404, 409, 500), debounced live search (300ms), and real-time DevTools Network inspection."
    )

    output_path = os.path.join(base_dir, "25CS052_Prac13.docx")
    doc.save(output_path)
    print(f"Successfully generated laboratory report: {output_path}")

if __name__ == "__main__":
    create_document()
