import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Page Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Header / College Info
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("CHANDUBHAI S. PATEL INSTITUTE OF TECHNOLOGY (CSPIT)\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    run_inst.bold = True
    run_inst.font.size = Pt(13)
    run_inst.font.color.rgb = RGBColor(30, 41, 59)

    run_sub = title_p.add_run("ITUE203 - Web Development Frameworks / Database Systems\n")
    run_sub.bold = True
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(37, 99, 235)

    run_prac = title_p.add_run("PRACTICAL - 15 FINAL LABORATORY SUBMISSION REPORT\n")
    run_prac.bold = True
    run_prac.font.size = Pt(15)
    run_prac.font.color.rgb = RGBColor(15, 23, 42)

    run_theme = title_p.add_run("Final Integration, Testing, Documentation, GitHub Submission, and Local Deployment\n")
    run_theme.italic = True
    run_theme.font.size = Pt(11)
    run_theme.bold = True

    # Student metadata table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_table.rows[0].cells[0].text = "Student Name: Jay Navadiya"
    meta_table.rows[0].cells[1].text = "Student ID / Roll No: 25CS052"
    meta_table.rows[1].cells[0].text = "Degree / Semester: B.Tech CSE (Div-I)"
    meta_table.rows[1].cells[1].text = "Date of Submission: October 21, 2026"

    for row in meta_table.rows:
        for cell in row.cells:
            set_cell_background(cell, "F1F5F9")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.bold = True
                    run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Problem Definition & Syllabus Mapping
    h1 = doc.add_heading("1. Aim, Objectives & Problem Definition", level=1)
    h1.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    p_aim = doc.add_paragraph()
    p_aim.add_run("AIM: ").bold = True
    p_aim.add_run(
        "To integrate all modules into a single deployable StudentHub portal. Verify navigation, responsiveness, "
        "authentication, database operations, security checks, AJAX functionality, and documentation. Deploy locally using "
        "XAMPP/Apache virtual host and submit the finalized GitHub repository."
    )

    doc.add_heading("1.1 Course Outcome (CO) Mapping", level=2)
    p_co = doc.add_paragraph()
    p_co.add_run("• CO2: ").bold = True
    p_co.add_run("Build dynamic, secure, and accessible database-driven web applications using relational MySQL schema and PHP.\n")
    p_co.add_run("• CO4: ").bold = True
    p_co.add_run("Implement hardened session statefulness, Role-Based Access Control (RBAC), and RESTful API endpoints.\n")
    p_co.add_run("• CO5: ").bold = True
    p_co.add_run("Deploy, document, and test full-stack web applications with local Apache Virtual Hosts and version-controlled Git workflows.")

    # Section 2: Submission Deliverables Summary
    h2 = doc.add_heading("2. Practical 15 Submission Deliverables Summary", level=1)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    deliv_table = doc.add_table(rows=6, cols=3)
    deliv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    deliv_headers = ["Item", "Submission Requirement", "Status & Artifact Link"]
    for i, h in enumerate(deliv_headers):
        deliv_table.rows[0].cells[i].text = h
        set_cell_background(deliv_table.rows[0].cells[i], "0F172A")
        p = deliv_table.rows[0].cells[i].paragraphs[0]
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].bold = True

    deliv_rows = [
        ("1", "GitHub Repo Link", "https://github.com/cs-cspit/Student-Portal-25CS052 (Integrated source code link)"),
        ("2", "Database Dump", "studenthub.sql (Full 3NF database dump with tables, FKs, and seed data)"),
        ("3", "README.md File", "README.md (Complete installation, testing, and deployment documentation)"),
        ("4", "Test & Security Report", "TEST_AND_SECURITY_REPORT.md (Comprehensive test matrix, security & accessibility audit)"),
        ("5", "Screenshots / Demo", "PRAC SS/PRACTICAL 15/ (16 high-quality light-themed verification captures)"),
    ]

    for r_idx, r_data in enumerate(deliv_rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = deliv_table.rows[r_idx].cells[c_idx]
            cell.text = val
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 3: Key Questions & Architectural Analysis
    h3 = doc.add_heading("3. Key Questions & Architectural Analysis", level=1)
    h3.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    q_items = [
        ("1. Are all pages properly linked and navigable?",
         "Yes. All portal modules (Home, Login, Register, Admin Dashboard, Student Dashboard, Student CRUD, Events, Timetable, Material, About, Contact) share a unified header and footer with dynamic role-adaptive navigation, persistent breadcrumbs, and zero broken links."),
        ("2. Are sessions, roles, validations, and database operations working end-to-end?",
         "Yes. Secure statefulness is guaranteed via session_regenerate_id(true), cross-tab epoch synchronization, and a 10-minute inactivity timer. RBAC strictly enforces role isolation. Relational MySQL transactions execute via prepared statements, ensuring 100% data integrity."),
        ("3. Are security and accessibility checklists completed?",
         "Yes. The platform passes all OWASP Top 10 web security checks (SQL injection defense, Bcrypt password hashing, session fixation defense, XSS escaping, RBAC 403 access control) and complies with WCAG 2.1 AA accessibility guidelines (15.8:1 light theme contrast ratio, visible focus rings, ARIA landmarks)."),
        ("4. Is project documentation sufficient for installation and testing?",
         "Yes. The comprehensive README.md provides turnkey instructions for PHP built-in server execution and Apache Virtual Host local deployment (http://studenthub.local), along with pre-configured credentials and testing steps."),
        ("5. Is the project repository clean and properly committed?",
         "Yes. The Git repository is cleanly organized with appropriate .gitignore rules, detailed semantic commit messages, and synchronized with origin/main on GitHub.")
    ]

    for q, ans in q_items:
        p = doc.add_paragraph()
        run_q = p.add_run(f"• {q}\n")
        run_q.bold = True
        run_q.font.color.rgb = RGBColor(15, 23, 42)
        run_ans = p.add_run(ans)
        run_ans.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(4)

    # Section 4: Intermediate & Advanced Deployment Extensions
    h4 = doc.add_heading("4. Deployment Extensions (Apache Virtual Host & Cloud)", level=1)
    h4.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    p_ext = doc.add_paragraph()
    p_ext.add_run("Intermediate Extension: Local Deployment with Apache Virtual Hosts\n").bold = True
    p_ext.add_run(
        "A customized VirtualHost configuration (apache_studenthub_vhost.conf) binds the domain 'studenthub.local' "
        "to the project root directory. With local DNS mapped in /etc/hosts (127.0.0.1 studenthub.local) and mod_rewrite "
        "activated, the application runs seamlessly under standard production web server conditions.\n\n"
    )
    p_ext.add_run("Advanced Extension: Online Cloud Deployment Strategy\n").bold = True
    p_ext.add_run(
        "The application architecture supports decoupling the static frontend onto edge content delivery networks (Vercel, "
        "Netlify) while containerizing the PHP backend and relational MySQL instance on platform-as-a-service providers "
        "(Railway, Render, AWS RDS) with environment variable isolation."
    )

    # Section 5: Verification Screenshots
    h5 = doc.add_heading("5. Visual Verification Screenshots (Light Theme Mode)", level=1)
    h5.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    screenshots_dir = os.path.join(os.path.dirname(__file__), "PRAC SS", "Practical 15")
    
    screenshot_items = [
        ("01_integrated_home_landing_page.png", "Figure 15.1: Integrated StudentHub Portal Homepage (Light Theme) with Brand Navigation and Module Cards"),
        ("02_user_login_auth_gateway.png", "Figure 15.2: Authentication Gateway (login.php) with Role Switching Tabs, Demo Credential Pills & Remember-Me"),
        ("03_user_registration_validation.png", "Figure 15.3: Student Self-Registration Interface (register.php) with Dual-Layer Field Validation"),
        ("04_admin_dashboard_full_overview.png", "Figure 15.4: Executive Admin Dashboard with Dynamic KPIs, Registrations Feed & Monthly Analytics Chart"),
        ("05_admin_audit_logs_security_trail.png", "Figure 15.5: Relational Audit Trail (audit_logs) Recording Administrative Actions and Security Violations"),
        ("06_student_dashboard_overview.png", "Figure 15.6: Personalized Student Academic Dashboard (student_dashboard.php) Logged in as Jay Navadiya"),
        ("07_student_crud_management_interface.png", "Figure 15.7: Student Management Directory (students.php) with 3NF Relational Records and Department Badges"),
        ("08_ajax_student_add_modal.png", "Figure 15.8: Asynchronous AJAX Add Student Modal Dialog with Real-Time Input Validation"),
        ("09_ajax_search_and_filter_live.png", "Figure 15.9: Live AJAX Search and Department Filtering without Full Page Reload"),
        ("10_events_catalog_and_registrations.png", "Figure 15.10: Campus Events Catalog and Relational Registration Junction (events.php)"),
        ("11_security_10min_session_timeout_modal.png", "Figure 15.11: 10-Minute Zero-Drift Session Inactivity Warning Modal with Live Countdown"),
        ("12_security_rbac_403_access_denied.png", "Figure 15.12: Strict Role-Based Access Control Interception (HTTP 403 Forbidden Access Denied Alert)"),
        ("13_responsive_mobile_viewport.png", "Figure 15.13: Mobile Responsive Viewport Layout (iPhone 14 Pro) Demonstrating Responsive Grid"),
        ("14_apache_virtual_host_configuration.png", "Figure 15.14: Apache Virtual Host Local Deployment Configuration (studenthub.local / httpd-vhosts.conf)"),
        ("15_mysql_studenthub_database_schema_dump.png", "Figure 15.15: MySQL studenthub Database Dump & Verification with Table Counts and Foreign Keys"),
        ("16_git_clean_repository_and_commit_history.png", "Figure 15.16: Clean Git Status, Origin Remote Link & Semantic Commit History Verification"),
    ]

    for fname, caption in screenshot_items:
        fpath = os.path.join(screenshots_dir, fname)
        if os.path.exists(fpath):
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            run_cap = p_cap.add_run(caption)
            run_cap.bold = True
            run_cap.font.size = Pt(10)
            run_cap.font.color.rgb = RGBColor(30, 41, 59)

            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_after = Pt(10)
            try:
                # Add picture scaled to 6.2 inches wide
                doc.add_picture(fpath, width=Inches(6.2))
            except Exception as e:
                p_img.add_run(f"[Error inserting image {fname}: {e}]")
        else:
            p_err = doc.add_paragraph(f"[Image file not found: {fname}]")
            p_err.runs[0].font.color.rgb = RGBColor(220, 38, 38)

    # Section 6: Evaluation Strategy & Viva Q&A
    h6 = doc.add_heading("6. Evaluation Strategy & Viva Questions and Answers", level=1)
    h6.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    viva_qa = [
        ("Q1: What is the purpose of session_regenerate_id(true) during login?",
         "Answer: It prevents Session Fixation attacks. If an attacker tricks a user into using a known session ID, regenerating the session ID upon authentication issues a new cryptographic token and invalidates the old one, destroying any hijacked pre-authentication session."),
        ("Q2: How does the application enforce Role-Based Access Control (RBAC)?",
         "Answer: Every privileged PHP script validates $_SESSION['role']. If a student attempts to access an administrative script (like dashboard.php), the request is intercepted, an HTTP 403 status code is sent, a security violation is logged in `audit_logs`, and an Access Denied view is rendered."),
        ("Q3: How does the 10-minute zero-drift session timeout work across multiple browser tabs?",
         "Answer: It stores the target expiration timestamp (Date.now() + 600000) in localStorage ('campusly_session_expire_at') and listens to 'storage' events. When a user interacts in one tab, the epoch timestamp updates across all open tabs, preventing false logouts and tab-throttling drift."),
        ("Q4: Why are prepared statements essential for all database queries?",
         "Answer: Prepared statements separate the SQL code from user-supplied data. Parameters are sent out-of-band and treated strictly as literals, completely neutralizing SQL Injection attacks regardless of user input characters."),
        ("Q5: How is local deployment configured with Apache Virtual Hosts?",
         "Answer: An Apache VirtualHost directive specifies ServerName studenthub.local and DocumentRoot pointing to the application directory. Mapping 127.0.0.1 studenthub.local in /etc/hosts routes browser requests through Apache with full mod_rewrite and .htaccess support.")
    ]

    for q, a in viva_qa:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        rq = p.add_run(f"• {q}\n")
        rq.bold = True
        rq.font.color.rgb = RGBColor(15, 23, 42)
        ra = p.add_run(a)
        ra.font.size = Pt(9.5)

    output_filename = "25CS052_Prac15.docx"
    doc.save(output_filename)
    print(f"✅ Document successfully created: {output_filename}")

if __name__ == "__main__":
    create_document()
