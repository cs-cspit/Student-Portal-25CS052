import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Set page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Header / College Info
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("CHANDUBHAI S. PATEL INSTITUTE OF TECHNOLOGY (CSPIT)\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    run_inst.bold = True
    run_inst.font.size = Pt(13)
    run_inst.font.color.rgb = RGBColor(30, 41, 59)

    run_sub = title_p.add_run("ITUE203 - Web Development Frameworks / Information Systems Security\n")
    run_sub.bold = True
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(37, 99, 235)

    run_prac = title_p.add_run("PRACTICAL - 10 LABORATORY REPORT\n")
    run_prac.bold = True
    run_prac.font.size = Pt(14)
    run_prac.font.color.rgb = RGBColor(15, 23, 42)

    run_theme = title_p.add_run("Secure Login/Logout with Sessions, Role-Based Access Control (RBAC), Session Fixation Defense & Timeout Handling")
    run_theme.italic = True
    run_theme.font.size = Pt(10)

    # Student metadata table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_table.rows[0].cells[0].text = "Student Name: Jay Navadiya"
    meta_table.rows[0].cells[1].text = "Student ID / Roll No: 25CS052"
    meta_table.rows[1].cells[0].text = "Degree / Semester: B.Tech CSE (Div-I)"
    meta_table.rows[1].cells[1].text = "Submission Date: October 14, 2026"

    for row in meta_table.rows:
        for cell in row.cells:
            set_cell_background(cell, "F1F5F9")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.bold = True
                    run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Aim & Objectives
    h1 = doc.add_heading("1. Aim & Practical Objectives", level=1)
    h1.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    p_aim = doc.add_paragraph()
    r_aim_bold = p_aim.add_run("AIM: ")
    r_aim_bold.bold = True
    p_aim.add_run(
        "To develop a production-grade, secure authentication and session management architecture using PHP and MySQLi, "
        "incorporating bcrypt password verification (password_verify()), session fixation prevention via session ID regeneration "
        "(session_regenerate_id(true)), strict Role-Based Access Control (RBAC) separating Student and Admin privileges, and "
        "an enforced 10-minute idle session timeout mechanism."
    )

    p_obj = doc.add_paragraph()
    p_obj.add_run("Key Deliverables & Specifications:").bold = True
    objectives = [
        ("db.php", "MySQL database connection using object-oriented MySQLi."),
        ("login.php", "User-friendly, white-themed login interface supporting email and password authentication."),
        ("authenticate.php", "Backend credential validation using password_verify() and session_regenerate_id(true)."),
        ("dashboard.php", "Main central portal dashboard featuring session verification, timeout display, and fixation proof."),
        ("admin_dashboard.php", "Administrative control center restricted strictly to users with the 'admin' role."),
        ("student_dashboard.php", "Student academic portal restricted strictly to users with the 'student' role."),
        ("logout.php", "Comprehensive session teardown (clearing $_SESSION, invalidating cookie, and invoking session_destroy())."),
        ("timeout.php", "Idle session timeout tracker automatically expiring sessions after 10 minutes (600 seconds)."),
        ("style.css", "Modern, responsive, clean white-themed CSS stylesheet."),
        ("user_registration.sql", "Complete MySQL schema dump with users table and seeded test accounts with bcrypt hashes.")
    ]
    for filename, desc in objectives:
        p_item = doc.add_paragraph(style='List Bullet')
        r_f = p_item.add_run(f"{filename}: ")
        r_f.bold = True
        p_item.add_run(desc)

    # Section 2: Security Concepts & Technical Explanation
    h2 = doc.add_heading("2. Security Architecture & Theoretical Concepts", level=1)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    doc.add_heading("2.1 Password Hashing & Verification (password_verify)", level=2)
    p_sec1 = doc.add_paragraph(
        "Plaintext passwords are never stored in the database. When credentials are saved, PHP's password_hash() "
        "creates a salted, slow crypt-hash using the Blowfish (bcrypt) algorithm ($2y$). During authentication, "
        "password_verify() verifies the submitted plaintext password against the stored bcrypt hash in constant time, "
        "preventing timing attacks and rainbow-table compromises."
    )

    doc.add_heading("2.2 Session Fixation Defense (session_regenerate_id(true))", level=2)
    p_sec2 = doc.add_paragraph(
        "Session Fixation occurs when an attacker forces a victim to use a known session identifier (e.g. via a manipulated URL "
        "or pre-set cookie). If the server continues using that same session identifier after authentication, the attacker can hijack "
        "the account. To prevent this vulnerability, authenticate.php immediately calls session_regenerate_id(true) upon valid login. "
        "This destroys the pre-login session file and issues a completely new, cryptographically random Session ID."
    )

    doc.add_heading("2.3 Role-Based Access Control (RBAC)", level=2)
    p_sec3 = doc.add_paragraph(
        "Users are assigned distinct roles ('admin' or 'student'). Each privileged route validates the active session role: "
        "admin_dashboard.php intercepts students and displays an explicit '403 Unauthorized Access' notice, while student_dashboard.php "
        "similarly intercepts administrator accounts attempting unauthorized entry."
    )

    doc.add_heading("2.4 Inactivity Session Timeout (10 Minutes / 600s)", level=2)
    p_sec4 = doc.add_paragraph(
        "Session abandonment poses significant risks on shared workstations. timeout.php monitors $_SESSION['last_activity']. "
        "If the time delta exceeds 600 seconds (10 minutes), the session is terminated and the user is redirected to login.php?msg=timeout."
    )

    # Section 3: Complete Source Code
    h3 = doc.add_heading("3. Complete Source Code Implementation", level=1)
    h3.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    files_to_include = [
        ("db.php", "MySQL Database Connection using MySQLi", "/Users/jaynavadiya07/Desktop/Project/db.php"),
        ("timeout.php", "Session Inactivity Timeout Handling (10 Minutes)", "/Users/jaynavadiya07/Desktop/Project/timeout.php"),
        ("login.php", "White-Themed Login Interface Form", "/Users/jaynavadiya07/Desktop/Project/login.php"),
        ("authenticate.php", "Authentication, password_verify() & session_regenerate_id()", "/Users/jaynavadiya07/Desktop/Project/authenticate.php"),
        ("dashboard.php", "Main Central Dashboard with Session Fixation Verification", "/Users/jaynavadiya07/Desktop/Project/dashboard.php"),
        ("student_dashboard.php", "Student Portal with RBAC Access Guard", "/Users/jaynavadiya07/Desktop/Project/student_dashboard.php"),
        ("admin_dashboard.php", "Administrator Control Center with RBAC Access Guard", "/Users/jaynavadiya07/Desktop/Project/admin_dashboard.php"),
        ("logout.php", "Secure Logout and Complete Session Destruction", "/Users/jaynavadiya07/Desktop/Project/logout.php"),
        ("style.css", "Clean White-Themed CSS Stylesheet", "/Users/jaynavadiya07/Desktop/Project/style.css"),
        ("user_registration.sql", "Exported MySQL Database Schema & Seed Data", "/Users/jaynavadiya07/Desktop/Project/user_registration.sql")
    ]

    for fname, fdesc, fpath in files_to_include:
        h_code = doc.add_heading(f"3.{files_to_include.index((fname, fdesc, fpath))+1} File: {fname} – {fdesc}", level=2)
        h_code.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                code_text = f.read()
        else:
            code_text = "// File not found"

        # Add code in shaded table/paragraph
        code_p = doc.add_paragraph()
        code_p.paragraph_format.space_before = Pt(4)
        code_p.paragraph_format.space_after = Pt(8)
        run_code = code_p.add_run(code_text)
        run_code.font.name = 'Courier New'
        run_code.font.size = Pt(8.5)
        run_code.font.color.rgb = RGBColor(30, 41, 59)

    # Section 4: Test Report & Execution Screenshots
    h4 = doc.add_heading("4. Test Report & Verification Screenshots", level=1)
    h4.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    p_test_intro = doc.add_paragraph(
        "The following test cases demonstrate the full end-to-end functionality, security controls, RBAC enforcement, "
        "and session lifecycle in clean white theme as required:"
    )

    test_screenshots = [
        ("01_student_login_success.png", "Test Case 1: Successful Student Login", 
         "Student enters valid credentials (jay@campusly.com / Student@123) and is successfully authenticated and redirected to student_dashboard.php."),
        
        ("02_admin_login_success.png", "Test Case 2: Successful Admin Login", 
         "Administrator enters valid credentials (admin@campusly.com / Admin@123) and is redirected to admin_dashboard.php showing live database user records."),
        
        ("03_invalid_email_password_message.png", "Test Case 3: Invalid Email/Password Error Notification", 
         "User submits an incorrect password; authenticate.php rejects the request and displays the red alert 'Invalid email address or password. Please verify your credentials and try again.'"),
        
        ("04_student_unauthorized_admin_dashboard.png", "Test Case 4: Student Attempting to Access Admin Dashboard (Unauthorized Access)", 
         "A logged-in Student attempts to navigate to admin_dashboard.php; the RBAC guard intercepts the request and renders a prominent '403 - Unauthorized Access Denied' notice."),
        
        ("05_admin_unauthorized_student_dashboard.png", "Test Case 5: Admin Attempting to Access Student Dashboard (Unauthorized Access)", 
         "A logged-in Administrator attempts to access student_dashboard.php; the RBAC guard detects role mismatch and renders an explicit '403 - Unauthorized Access Denied' notice."),
        
        ("06_successful_logout.png", "Test Case 6: Successful Secure Logout", 
         "Clicking 'Sign Out' invokes logout.php, unsets all session variables, deletes the session cookie, destroys server session state, and redirects to login.php with green confirmation."),
        
        ("07_dashboard_access_denied_after_logout.png", "Test Case 7: Dashboard Access Denied After Logout", 
         "Attempting to access dashboard.php or privileged routes after session destruction redirects immediately to login.php?msg=unauthorized with an amber access denied banner."),
        
        ("08_session_timeout_redirection.png", "Test Case 8: Session Timeout and Redirection to Login", 
         "After exceeding the 10-minute inactivity limit (600 seconds), timeout.php invalidates the session and redirects the user to login.php?msg=timeout with an expiration alert."),
        
        ("09_session_id_regeneration_verification.png", "Test Case 9: Session ID Regeneration After Successful Login (Fixation Defense)", 
         "Verification table confirming session_regenerate_id(true) changed the initial pre-login session ID to a new crypto session ID upon authentication.")
    ]

    ss_dir = "/Users/jaynavadiya07/Desktop/Project/PRAC SS/Practical 10"

    for img_name, test_title, test_obs in test_screenshots:
        img_path = os.path.join(ss_dir, img_name)
        
        h_t = doc.add_heading(test_title, level=2)
        h_t.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        
        p_desc = doc.add_paragraph()
        p_desc.add_run("Observation & Verification: ").bold = True
        p_desc.add_run(test_obs)

        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(6.2))
            
            p_caption = doc.add_paragraph()
            p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cap = p_caption.add_run(f"Figure: {test_title}")
            r_cap.font.size = Pt(8.5)
            r_cap.italic = True
            r_cap.font.color.rgb = RGBColor(100, 116, 139)
        else:
            p_missing = doc.add_paragraph(f"[Image file not found: {img_name}]")
            p_missing.runs[0].font.color.rgb = RGBColor(220, 38, 38)

    # Section 5: Conclusion
    h5 = doc.add_heading("5. Conclusion", level=1)
    h5.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph(
        "Practical 10 was executed and verified successfully. All core objectives were achieved: "
        "secure MySQLi database connection, bcrypt hashing with password_verify(), session fixation mitigation via "
        "session_regenerate_id(true), bidirectional Role-Based Access Control between student and administrator portals, "
        "inactivity timeout handling (10 minutes / 600 seconds), and secure session teardown on logout."
    )

    out_file = "/Users/jaynavadiya07/Desktop/Project/25CS052_Prac10.docx"
    doc.save(out_file)
    print(f"Document successfully created at: {out_file}")

    # Also save in OneDrive if directory exists
    onedrive_dir = "/Users/jaynavadiya07/Library/CloudStorage/OneDrive-Personal"
    if os.path.exists(onedrive_dir):
        od_file = os.path.join(onedrive_dir, "25CS052_Prac10.docx")
        doc.save(od_file)
        print(f"Document copy saved to OneDrive at: {od_file}")

if __name__ == "__main__":
    create_document()
