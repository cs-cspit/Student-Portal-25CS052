import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Header / College Info
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("CHANDUBHAI S. PATEL INSTITUTE OF TECHNOLOGY (CSPIT)\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    run_inst.bold = True
    run_inst.font.size = Pt(13)
    run_inst.font.color.rgb = RGBColor(30, 41, 59)

    run_sub = title_p.add_run("ITUE203 - Web Development Frameworks & Database Systems\n")
    run_sub.bold = True
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(37, 99, 235)

    run_prac = title_p.add_run("PRACTICAL - 11 LABORATORY TEST REPORT\n")
    run_prac.bold = True
    run_prac.font.size = Pt(14)
    run_prac.font.color.rgb = RGBColor(15, 23, 42)

    run_theme = title_p.add_run("Student Management CRUD with Search, Course Filtering, Soft Delete & CSV Export")
    run_theme.italic = True
    run_theme.font.size = Pt(10)

    # Student metadata table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_table.rows[0].cells[0].text = "Student Name: Jay Navadiya"
    meta_table.rows[0].cells[1].text = "Student ID / Roll No: 25CS052"
    meta_table.rows[1].cells[0].text = "Degree / Semester: B.Tech CSE (Div-I)"
    meta_table.rows[1].cells[1].text = "Submission Date: October 16, 2026"

    for row in meta_table.rows:
        for cell in row.cells:
            set_cell_background(cell, "F1F5F9")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.bold = True
                    run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Aim & Objectives
    h1 = doc.add_heading("1. Aim & Practical Objectives", level=1)
    h1.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    p_aim = doc.add_paragraph()
    p_aim.add_run("AIM: ").bold = True
    p_aim.add_run(
        "To develop a comprehensive, production-grade student management module (students.php) supporting full "
        "Create, Read, Update, and Delete (CRUD) capabilities using parameterized MySQLi statements, real-time "
        "search by name and email, course-wise department filtering, confirmation modal dialogs before delete execution, "
        "advanced soft deletion via an active/inactive status flag, and intermediate extension of CSV export opened in Excel."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Section 2: Test Report & Execution Screenshots
    h2 = doc.add_heading("2. Verified Test Report & Execution Screenshots", level=1)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    p_test_intro = doc.add_paragraph(
        "Every requirement outlined in Practical 11 instructions has been executed and captured in high-resolution white theme:"
    )

    test_screenshots = [
        ("01_student_records_table.png", "Requirement 1: Student records displayed in a table",
         "The catalog displays 20+ seeded student records in a clean tabular view, including ID, Student Name, University Email, Department, Status (Active badge), and Actions."),
        
        ("02_add_student_modal.png", "Requirement 2: Add Student Modal Interface",
         "Interactive modal form collecting student Full Name, University Email, Course/Department dropdown, and initial status with validation safeguards."),
        
        ("03_add_student_success.png", "Requirement 3: Successful addition of a new student",
         "Backend execution of parameterized MySQLi INSERT statement; displays green Operation Success banner and confirms creation of new student record."),
        
        ("04_search_functionality.png", "Requirement 4: Search functionality (Name & Email)",
         "Real-time search query executed via MySQLi parameterized LIKE statement, filtering records matching student names or email domains instantly."),
        
        ("05_course_wise_filtering.png", "Requirement 5: Course-wise filtering",
         "Departmental dropdown filter applied (e.g. IT); dynamically restricts displayed students to the selected academic branch."),
        
        ("06_edit_student_modal.png", "Requirement 6: Edit Student Record Modal",
         "Pre-populated edit modal allowing modification of student name, email, department, and active/inactive status."),
        
        ("07_edit_student_success.png", "Requirement 7: Successful update of student details",
         "Execution of MySQLi UPDATE query with duplicate email prevention; renders green success notification and updates catalog data."),
        
        ("08_delete_confirmation_message.png", "Requirement 8: Delete confirmation message",
         "Confirmation dialog displayed before execution, warning the administrator and verifying intent to soft delete/deactivate the student record."),
        
        ("09_successful_soft_delete.png", "Requirement 9: Successful soft delete (student status changed to inactive)",
         "Advanced Extension: Instead of irreversible hard deletion, updates student status to 'Inactive' with distinct visual badge and Reactivate action."),
        
        ("10_success_error_messages.png", "Requirement 10: Clear Success and Error feedback messages",
         "Demonstration of error handling when invalid or duplicate email credentials are submitted, rendering distinct red error feedback banner."),
        
        ("11_exported_csv_opened_in_excel.png", "Requirement 11: Exported CSV file opened in Excel",
         "Intermediate Extension: campusly_students.csv exported using PHP fputcsv() output stream and displayed opened in Microsoft Excel with formatted columns and active status.")
    ]

    ss_dir = "/Users/jaynavadiya07/Desktop/Project/PRAC SS/Practical 11"

    for img_name, test_title, test_obs in test_screenshots:
        img_path = os.path.join(ss_dir, img_name)
        h_t = doc.add_heading(test_title, level=2)
        h_t.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        p_obs = doc.add_paragraph()
        p_obs.add_run("Observation / Verification: ").bold = True
        p_obs.add_run(test_obs)

        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(12)
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(6.2))
        else:
            p_err = doc.add_paragraph(f"[Image file not found: {img_name}]")
            p_err.runs[0].font.color.rgb = RGBColor(220, 38, 38)

    output_path = "/Users/jaynavadiya07/Desktop/Project/25CS052_Prac11.docx"
    doc.save(output_path)
    print(f"Practical 11 Report successfully generated at {output_path}")

if __name__ == "__main__":
    create_document()
